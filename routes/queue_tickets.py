import os
from flask import Blueprint, render_template, request, jsonify, session, redirect, flash, url_for, current_app, send_file
from werkzeug.utils import secure_filename
from db_config import get_db_connection
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import io
from datetime import datetime, date

# 💡 កំណត់ Blueprint
queue_bp = Blueprint('queue', __name__)

# 💡 កំណត់ទីតាំង Folder សម្រាប់ទុក File Slideshow របស់ TV
TV_SLIDES_FOLDER = os.path.join('static', 'uploads', 'tv_slides')
# 💡 កំណត់ទីតាំង Folder សម្រាប់ទុក File សំឡេងកណ្តឹង
TV_SOUNDS_FOLDER = os.path.join('static', 'sounds')

# 💡 បង្កើត Folder ទាំងនេះដោយស្វ័យប្រវត្តិបើវាមិនទាន់មាន (ពេលកូដរត់)
os.makedirs(TV_SLIDES_FOLDER, exist_ok=True)
os.makedirs(TV_SOUNDS_FOLDER, exist_ok=True)

# ==========================================
# ផ្ទាំងបញ្ជរ (ទាញយកបញ្ជរពី Database មកបង្ហាញ)
# ==========================================
@queue_bp.route('/counter')
def counter_dashboard():
    if 'loggedin' not in session:
        return redirect(url_for('login')) 
    
    allowed_roles = ['Super Admin', 'Admin', 'Application Receiver']
    if session.get('role_name') not in allowed_roles:
        flash('សុំទោស! អ្នកមិនមានសិទ្ធិចូលប្រើប្រាស់ផ្ទាំងនេះទេ។', 'danger')
        return redirect(url_for('dashboard')) 
    
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    
    assigned_counter = None
    reject_reasons = [] 
    
    try:
        # 🌟 ស្វែងរកតែបញ្ជរណា ដែល Admin បានចាត់តាំងឱ្យ User ម្នាក់នេះ
        cursor.execute("""
            SELECT id, counter_name_kh, counter_name_en, status 
            FROM counters 
            WHERE current_user_id = %s LIMIT 1
        """, (session.get('id'),)) 
        assigned_counter = cursor.fetchone()
        
        # ទាញយកមូលហេតុបដិសេធ
        cursor.execute("SELECT * FROM ref_app_reject_reasons WHERE is_active = 1 ORDER BY id ASC")
        reject_reasons = cursor.fetchall()
    except Exception as e:
        print(f"❌ Error loading counter dashboard: {e}")
    finally:
        if cursor: cursor.close()
        if conn: conn.close()
    
    return render_template('counter_dashboard.html', 
                           full_name=session.get('full_name', 'Staff'),
                           assigned_counter=assigned_counter,
                           reject_reasons=reject_reasons)

# ==========================================
# API សម្រាប់ចុច "ហៅអ្នកបន្ទាប់" និង "បញ្ចប់ & សម្រាក"
# ==========================================
@queue_bp.route('/api/call_next', methods=['POST'])
def call_next():
    if 'loggedin' not in session:
        return jsonify({'status': 'error', 'message': 'Unauthorized'}), 401

    data = request.json
    counter_id = data.get('counter_no') 
    user_id = session.get('id')
    
    action_type = data.get('action_type', 'complete') 
    reject_reason = data.get('reject_reason', None)
    
    if not counter_id:
        return jsonify({'status': 'error', 'message': 'សូមជ្រើសរើសបញ្ជរ!'})

    # កំណត់ម៉ោងកម្ពុជា
    from datetime import datetime, timedelta
    today_date_khmer = (datetime.utcnow() + timedelta(hours=7)).strftime('%Y-%m-%d')

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    
    try:
        conn.start_transaction()
        cursor.execute("UPDATE counters SET status = 'Active', current_user_id = %s WHERE id = %s", (user_id, counter_id))
        
        # 🌟 ដំណាក់កាលទី ១៖ ឆែកមើលសំបុត្រដែលកំពុងដំណើរការ (Processing) របស់បញ្ជរនេះ
        cursor.execute("""
            SELECT qt.id, qt.application_no, qt.queue_number, c.name_en 
            FROM queue_tickets qt
            LEFT JOIN candidates c ON qt.application_no = c.application_no
            WHERE qt.counter_no = %s AND qt.status = 'Processing' AND DATE(qt.updated_at) = %s 
            LIMIT 1
        """, (counter_id, today_date_khmer))
        current_ticket = cursor.fetchone()

        # 🚀 ដំណោះស្រាយការពារការ Refresh (F5) និង Double-Click
        if current_ticket and action_type == 'call_next_new':
            # បើគាត់ចុចហៅលេខថ្មី តែនៅមានលេខចាស់មិនទាន់ចប់ យើងប្រគល់លេខចាស់នោះឱ្យគាត់វិញ (មិនឱ្យទៅ Scanner ទេ)
            conn.commit()
            ticket_data = {
                'queue_number': current_ticket['queue_number'],
                'candidate_name': current_ticket['name_en'] if current_ticket['name_en'] else 'មិនមានព័ត៌មាន',
                'exam_code': current_ticket['application_no'] if current_ticket['application_no'] else 'មិនមានព័ត៌មាន'
            }
            return jsonify({
                'status': 'success', 
                'message': f'សូមបន្តពិនិត្យលេខ {current_ticket["queue_number"]} ដែលនៅសេសសល់ឱ្យចប់សិន!', 
                'ticket': ticket_data
            })

        # ប្រសិនបើជាការចុច "បដិសេធ", "យល់ព្រម", ឬ "សម្រាក" លើសំបុត្រដែលកំពុងកាន់
        elif current_ticket:
            app_no = current_ticket['application_no']

            if action_type == 'reject':
                if not reject_reason or str(reject_reason).strip() == '':
                    raise ValueError("សូមបញ្ចូលមូលហេតុនៃការបដិសេធ!")
                    
                cursor.execute("UPDATE queue_tickets SET status = 'Rejected', reject_reason = %s, updated_at = NOW() WHERE id = %s", (str(reject_reason).strip(), current_ticket['id']))
                cursor.execute("UPDATE application_details SET status_id = 3, remark = %s WHERE application_no = %s", (str(reject_reason).strip(), app_no))
                cursor.execute("UPDATE candidates SET job_app_status = 'REJECTED' WHERE application_no = %s", (app_no,))

            elif action_type in ['complete', 'clear_only']: 
                # 🌟 កន្លែងនេះហើយដែលបញ្ជូនទិន្នន័យទៅ Scanner (ព្រោះ Status ដូរទៅ Completed)
                cursor.execute("UPDATE queue_tickets SET status = 'Completed', updated_at = NOW() WHERE id = %s", (current_ticket['id'],))
                
                # 🌟🌟 ចំណុចដែលត្រូវកែតម្រូវ៖ ត្រូវ Reset `scan_status` ឱ្យត្រឡប់ទៅជា 'Pending' វិញជានិច្ច 🌟🌟
                cursor.execute("UPDATE application_details SET scan_status = 'Pending', status_id = 1, remark = NULL WHERE application_no = %s", (app_no,))
        
        # 🌟 ដំណាក់កាលទី ២៖ ហៅអ្នកបន្ទាប់ ឬ គ្រាន់តែសម្អាតអេក្រង់
        if action_type == 'clear_only':
            conn.commit()
            return jsonify({'status': 'empty', 'message': 'បានបញ្ចប់សំបុត្រ និងសម្អាតអេក្រង់ជោគជ័យ!'})

        # ទាញយកសំបុត្របន្ទាប់
        cursor.execute("""
            SELECT qt.id, qt.queue_number, qt.application_no, c.name_en AS candidate_name
            FROM queue_tickets qt
            LEFT JOIN candidates c ON qt.application_no = c.application_no
            WHERE qt.status = 'Waiting' AND DATE(qt.created_at) = %s
            ORDER BY qt.created_at ASC LIMIT 1 FOR UPDATE
        """, (today_date_khmer,))
        next_ticket = cursor.fetchone()

        if next_ticket:
            cursor.execute("UPDATE queue_tickets SET status = 'Processing', counter_no = %s, processed_by = %s, updated_at = NOW() WHERE id = %s", (counter_id, user_id, next_ticket['id']))
            conn.commit()

            ticket_data = {
                'queue_number': next_ticket['queue_number'],
                'candidate_name': next_ticket['candidate_name'] if next_ticket['candidate_name'] else 'មិនមានព័ត៌មាន',
                'exam_code': next_ticket['application_no'] if next_ticket['application_no'] else 'មិនមានព័ត៌មាន'
            }
            return jsonify({'status': 'success', 'message': f'កំពុងហៅលេខ {next_ticket["queue_number"]}', 'ticket': ticket_data})
        else:
            conn.commit()
            return jsonify({'status': 'empty', 'message': 'មិនមានអ្នករង់ចាំទេ!'})
            
    except ValueError as ve:
        conn.rollback()
        return jsonify({'status': 'error', 'message': str(ve)})
    except Exception as e:
        conn.rollback()
        print(f"Error in call_next: {e}")
        return jsonify({'status': 'error', 'message': 'មានបញ្ហាប្រព័ន្ធ Server!'})
    finally:
        if cursor: cursor.close()
        if conn: conn.close()

   

from datetime import datetime, timedelta
from flask import session, request, flash, redirect, url_for, render_template, jsonify

# ==========================================
# ១. ផ្ទាំងគ្រប់គ្រងបញ្ជរ (Admin: Manage Counters)
# ==========================================
@queue_bp.route('/manage_counters', methods=['GET', 'POST'])
def manage_counters():
    if 'loggedin' not in session:
        return redirect(url_for('login'))
        
    allowed_roles = ['Admin', 'Super Admin', 'Center President', 'Head of Reception']
    if session.get('role_name') not in allowed_roles:
        flash('សុំទោស! អ្នកមិនមានសិទ្ធិចូលទំព័រគ្រប់គ្រងបញ្ជរនេះទេ។', 'danger')
        return redirect(url_for('dashboard'))

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    if request.method == 'POST':
        action = request.form.get('action')
        try:
            if action == 'add':
                name_kh = request.form.get('counter_name_kh').strip()
                name_en = request.form.get('counter_name_en').strip()
                cursor.execute("INSERT INTO counters (counter_name_kh, counter_name_en, status) VALUES (%s, %s, 'Available')", (name_kh, name_en))
                conn.commit()
                flash('បង្កើតបញ្ជរថ្មីបានជោគជ័យ!', 'success')
            elif action == 'edit':
                c_id = request.form.get('counter_id')
                name_kh = request.form.get('counter_name_kh').strip()
                name_en = request.form.get('counter_name_en').strip()
                cursor.execute("UPDATE counters SET counter_name_kh=%s, counter_name_en=%s WHERE id=%s", (name_kh, name_en, c_id))
                conn.commit()
                flash('កែប្រែបញ្ជរបានជោគជ័យ!', 'success')
            elif action == 'reset': 
                c_id = request.form.get('counter_id')
                cursor.execute("UPDATE counters SET status='Available', current_user_id=NULL WHERE id=%s", (c_id,))
                conn.commit()
                flash('បាន Reset បញ្ជរឱ្យទំនេរវិញជោគជ័យ!', 'success')
            elif action == 'delete':
                c_id = request.form.get('counter_id')
                cursor.execute("DELETE FROM counters WHERE id=%s", (c_id,))
                conn.commit()
                flash('លុបបញ្ជរបានជោគជ័យ!', 'success')
            elif action == 'assign':
                c_id = request.form.get('counter_id')
                user_id = request.form.get('user_id')
                if user_id:
                    cursor.execute("UPDATE counters SET current_user_id=%s, status='Available' WHERE id=%s", (user_id, c_id))
                    flash('បានចាត់តាំងមន្ត្រីចូលបញ្ជរជោគជ័យ!', 'success')
                else:
                    cursor.execute("UPDATE counters SET current_user_id=NULL, status='Available' WHERE id=%s", (c_id,))
                    flash('បានដកមន្ត្រីចេញពីបញ្ជរជោគជ័យ!', 'info')
                conn.commit()
        except Exception as e:
            conn.rollback()
            flash(f'មិនអាចអនុវត្តបានទេ៖ {str(e)}', 'danger')
        finally:
            cursor.close()
            conn.close()
            
        return redirect(url_for('queue.manage_counters'))

    try:
        today_date_khmer = (datetime.utcnow() + timedelta(hours=7)).strftime('%Y-%m-%d')

        cursor.execute("""
            SELECT u.id, u.full_name, r.name as role_name
            FROM users u
            LEFT JOIN roles r ON u.role_id = r.role_id
            WHERE u.is_active = 1
        """)
        officers_list = cursor.fetchall()

        # 🌟 ទី១៖ រាប់តែជោគជ័យ (Completed) សម្រាប់ទិន្នន័យសរុបខាងលើ
        cursor.execute("""
            SELECT 
                COUNT(q.id) as total_today,
                SUM(CASE WHEN c.gender IN ('F', 'FEMALE', 'ស្រី') THEN 1 ELSE 0 END) as female_today
            FROM queue_tickets q
            LEFT JOIN candidates c ON q.application_no = c.application_no
            WHERE DATE(q.updated_at) = %s
              AND q.status = 'Completed'
        """, (today_date_khmer,))
        stats = cursor.fetchone()
        total_today = stats['total_today'] if stats and stats['total_today'] else 0
        female_today = int(stats['female_today']) if stats and stats['female_today'] else 0

        # 🌟 ទី២៖ រាប់តែជោគជ័យ (Completed) សម្រាប់ទិន្នន័យបញ្ជរនីមួយៗ
        cursor.execute("""
            SELECT c.*, u.full_name as current_user_name,
            (
                SELECT COUNT(q.id) FROM queue_tickets q 
                WHERE q.counter_no = c.id 
                  AND DATE(q.updated_at) = %s 
                  AND q.status = 'Completed'
            ) as processed_today,
            COALESCE((
                SELECT SUM(CASE WHEN cand.gender IN ('F', 'FEMALE', 'ស្រី') THEN 1 ELSE 0 END)
                FROM queue_tickets q
                LEFT JOIN candidates cand ON q.application_no = cand.application_no
                WHERE q.counter_no = c.id 
                  AND DATE(q.updated_at) = %s 
                  AND q.status = 'Completed'
            ), 0) as female_processed_today
            FROM counters c 
            LEFT JOIN users u ON c.current_user_id = u.id
            ORDER BY c.id ASC
        """, (today_date_khmer, today_date_khmer))
        counters_list = cursor.fetchall()
        
    except Exception as e:
        counters_list, officers_list = [], []
        total_today, female_today = 0, 0
        print(f"Error loading manage_counters: {e}")
    finally:
        cursor.close()
        conn.close()

    return render_template('manage_counters.html', counters=counters_list, officers=officers_list, total_today=total_today, female_today=female_today)


# ==========================================
# ២. API សម្រាប់ Auto-Refresh ផ្ទាំងគ្រប់គ្រងបញ្ជរ (Silent Refresh)
# ==========================================
@queue_bp.route('/api/manage_counters_data', methods=['GET'])
def get_manage_counters_data():
    if 'loggedin' not in session:
        return jsonify({'status': 'error', 'message': 'Unauthorized'}), 401

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        today_date_khmer = (datetime.utcnow() + timedelta(hours=7)).strftime('%Y-%m-%d')

        # 🌟 ទី៣៖ រាប់តែជោគជ័យក្នុង API Refresh សម្រាប់ខាងលើ
        cursor.execute("""
            SELECT 
                COUNT(q.id) as total_today,
                SUM(CASE WHEN c.gender IN ('F', 'FEMALE', 'ស្រី') THEN 1 ELSE 0 END) as female_today
            FROM queue_tickets q
            LEFT JOIN candidates c ON q.application_no = c.application_no
            WHERE DATE(q.updated_at) = %s
              AND q.status = 'Completed'
        """, (today_date_khmer,))
        stats = cursor.fetchone()
        total_today = stats['total_today'] if stats and stats['total_today'] else 0
        female_today = int(stats['female_today']) if stats and stats['female_today'] else 0

        # 🌟 ទី៤៖ រាប់តែជោគជ័យក្នុង API Refresh សម្រាប់បញ្ជរនីមួយៗ
        cursor.execute("""
            SELECT c.id, c.status, u.full_name as current_user_name,
            (
                SELECT COUNT(q.id) FROM queue_tickets q 
                WHERE q.counter_no = c.id AND DATE(q.updated_at) = %s AND q.status = 'Completed'
            ) as processed_today,
            COALESCE((
                SELECT SUM(CASE WHEN cand.gender IN ('F', 'FEMALE', 'ស្រី') THEN 1 ELSE 0 END)
                FROM queue_tickets q
                LEFT JOIN candidates cand ON q.application_no = cand.application_no
                WHERE q.counter_no = c.id AND DATE(q.updated_at) = %s AND q.status = 'Completed'
            ), 0) as female_processed_today
            FROM counters c 
            LEFT JOIN users u ON c.current_user_id = u.id
            ORDER BY c.id ASC
        """, (today_date_khmer, today_date_khmer))
        counters_list = cursor.fetchall()
        
        active_count = sum(1 for c in counters_list if c['status'] == 'Active')

        return jsonify({
            'status': 'success',
            'total_today': total_today,
            'female_today': female_today,
            'active_count': active_count,
            'counters': counters_list
        })
    except Exception as e:
        return jsonify({'status': 'error', 'message': str(e)})
    finally:
        cursor.close()
        conn.close()

import io
import openpyxl
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime, date, timedelta
from flask import session, request, flash, redirect, url_for, send_file

# ==========================================
# ទាញយករបាយការណ៍មន្ត្រី (Officer Report)
# ==========================================
@queue_bp.route('/download_officer_report', methods=['POST'])
def download_officer_report():
    if 'loggedin' not in session: return redirect(url_for('login'))
    
    # 🌟 កំណត់ម៉ោងកម្ពុជាសម្រាប់ពេលបច្ចុប្បន្ន
    cambodia_now = datetime.utcnow() + timedelta(hours=7)
    
    start_date = request.form.get('start_date') or cambodia_now.strftime('%Y-%m-%d')
    end_date = request.form.get('end_date') or start_date
    current_datetime = cambodia_now.strftime('%d-%m-%Y %H:%M:%S')
    exporter_name = session.get('full_name', 'Admin')

    from app import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    
    try:
        # 🌟 ប្រើ DATE_ADD ដើម្បីបូក ៧ ម៉ោងទៅលើ q.updated_at ឱ្យក្លាយជាម៉ោងកម្ពុជា
        cursor.execute("""
            SELECT 
                DATE(DATE_ADD(q.updated_at, INTERVAL 7 HOUR)) as working_date, 
                u.full_name as officer_name, r.name as role_name,
                COUNT(q.id) as total_received,
                SUM(CASE WHEN q.status = 'Completed' THEN 1 ELSE 0 END) as total_approved,
                SUM(CASE WHEN q.status = 'Rejected' THEN 1 ELSE 0 END) as total_rejected
            FROM queue_tickets q
            JOIN users u ON q.processed_by = u.id LEFT JOIN roles r ON u.role_id = r.role_id
            WHERE q.processed_by IS NOT NULL 
              AND DATE(DATE_ADD(q.updated_at, INTERVAL 7 HOUR)) BETWEEN %s AND %s 
              AND q.status IN ('Completed', 'Rejected')
            GROUP BY DATE(DATE_ADD(q.updated_at, INTERVAL 7 HOUR)), u.id, u.full_name, r.name 
            ORDER BY working_date DESC, total_received DESC
        """, (start_date, end_date))
        summary_data = cursor.fetchall()

        cursor.execute("""
            SELECT 
                DATE(DATE_ADD(q.updated_at, INTERVAL 7 HOUR)) as working_date, 
                TIME(DATE_ADD(q.updated_at, INTERVAL 7 HOUR)) as working_time, 
                u.full_name as officer_name, r.name as role_name,
                q.counter_no, q.queue_number, c.application_no, c.name_en, a.last_name_kh, a.first_name_kh, 
                q.status as action_status, q.reject_reason
            FROM queue_tickets q
            JOIN users u ON q.processed_by = u.id LEFT JOIN roles r ON u.role_id = r.role_id
            LEFT JOIN candidates c ON q.application_no = c.application_no 
            LEFT JOIN application_details a ON c.application_no = a.application_no
            WHERE q.processed_by IS NOT NULL 
              AND DATE(DATE_ADD(q.updated_at, INTERVAL 7 HOUR)) BETWEEN %s AND %s 
              AND q.status IN ('Completed', 'Rejected')
            ORDER BY working_date DESC, u.full_name ASC, DATE_ADD(q.updated_at, INTERVAL 7 HOUR) ASC
        """, (start_date, end_date))
        detail_data = cursor.fetchall()

        wb = openpyxl.Workbook()

        def apply_standard_styles(ws):
            font_title = Font(name='Khmer OS Siemreap', size=15, bold=True, color="0f172a")
            font_subtitle = Font(name='Khmer OS Siemreap', size=12, color="475569")
            font_header = Font(name='Khmer OS Siemreap', size=13, bold=True, color="ffffff")
            font_data = Font(name='Khmer OS Siemreap', size=11)
            
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            align_center = Alignment(horizontal='center', vertical='center', wrap_text=False)
            align_left = Alignment(horizontal='left', vertical='center', wrap_text=False)
            header_fill = PatternFill(start_color="1e293b", end_color="1e293b", fill_type="solid")

            ws.row_dimensions[1].height = 35 
            ws.row_dimensions[2].height = 25 
            ws.row_dimensions[3].height = 25
            ws.row_dimensions[4].height = 35 

            for row_idx, row in enumerate(ws.iter_rows(), start=1):
                if row_idx > 4:
                    ws.row_dimensions[row_idx].height = 28
                    
                for cell in row:
                    if row_idx == 1: 
                        cell.font = font_title
                        cell.alignment = align_center
                    elif row_idx in [2, 3]: 
                        cell.font = font_subtitle
                        cell.alignment = align_center
                    elif row_idx == 4: 
                        cell.font = font_header
                        cell.alignment = align_center
                        cell.border = thin_border
                        cell.fill = header_fill
                    else: 
                        cell.font = font_data
                        cell.border = thin_border
                        if cell.column == 1 or cell.column == 2 or cell.column == 8: 
                            cell.alignment = align_center
                        else:
                            cell.alignment = align_left

            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                
                for cell in col:
                    if cell.row < 4: continue
                    try:
                        if cell.value:
                            cell_length = len(str(cell.value))
                            if cell_length > max_length: max_length = cell_length
                    except: pass
                
                if column == 'A':
                    ws.column_dimensions[column].width = 6 
                else:
                    adjusted_width = (max_length * 1.3) + 5
                    ws.column_dimensions[column].width = min(max(adjusted_width, 15), 60)

            max_col_letter = get_column_letter(ws.max_column)
            ws.merge_cells(f'A1:{max_col_letter}1')
            ws.merge_cells(f'A2:{max_col_letter}2')
            ws.merge_cells(f'A3:{max_col_letter}3')

        ws_summary = wb.active
        ws_summary.title = "Performance Summary"
        
        ws_summary.append([f"របាយការណ៍សង្ខេបមន្ត្រីទទួលពាក្យ (ពី {start_date} ដល់ {end_date})"])
        ws_summary.append([f"📅 ថ្ងៃទាញយក៖ {current_datetime} | 👤 ដោយ៖ {exporter_name}"])
        ws_summary.append([]) 
        
        headers1 = ["ល.រ", "កាលបរិច្ឆេទធ្វើការ", "ឈ្មោះមន្ត្រី", "តួនាទី", "សរុប (ទទួល)", "យល់ព្រម (Approve)", "បដិសេធ (Reject)"]
        ws_summary.append(headers1)
        
        for i, row in enumerate(summary_data, 1): 
            # ការពារការ Error ប្រសិនបើ row['working_date'] ជា String ស្រាប់
            w_date = row['working_date'].strftime('%Y-%m-%d') if isinstance(row['working_date'], date) else row['working_date']
            ws_summary.append([i, w_date, row['officer_name'], row['role_name'] or '-', row['total_received'] or 0, row['total_approved'] or 0, row['total_rejected'] or 0])

        apply_standard_styles(ws_summary) 

        ws_details = wb.create_sheet(title="Detailed List")
        
        ws_details.append([f"បញ្ជីបេក្ខជនលម្អិតតាមមន្ត្រី (ពី {start_date} ដល់ {end_date})"])
        ws_details.append([f"📅 ថ្ងៃទាញយក៖ {current_datetime} | 👤 ដោយ៖ {exporter_name}"])
        ws_details.append([]) 
        
        headers2 = ["ល.រ", "ថ្ងៃទី", "ម៉ោង", "ឈ្មោះមន្ត្រី", "តួនាទី", "លេខបញ្ជរ", "លេខរង់ចាំ", "លេខកូដប្រឡង", "ឈ្មោះខ្មែរ", "ឈ្មោះអង់គ្លេស", "ស្ថានភាព", "មូលហេតុបដិសេធ"]
        ws_details.append(headers2)
        
        for i, row in enumerate(detail_data, 1):
            full_kh = f"{row['last_name_kh'] or ''} {row['first_name_kh'] or ''}".strip()
            status_txt = "យល់ព្រម" if row['action_status'] == 'Completed' else "បដិសេធ"
            w_date = row['working_date'].strftime('%Y-%m-%d') if isinstance(row['working_date'], date) else row['working_date']
            
            ws_details.append([i, w_date, str(row['working_time']), row['officer_name'], row['role_name'] or '-', row['counter_no'] or '-', row['queue_number'] or '-', row['application_no'] or '-', full_kh, row['name_en'] or '-', status_txt, row['reject_reason'] or '-'])

        apply_standard_styles(ws_details) 

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=f"Officer_Report_{start_date}_to_{end_date}.xlsx", mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    except Exception as e:
        print(f"Error Report: {e}")
        flash(f'មានបញ្ហាក្នុងការទាញយករបាយការណ៍: {str(e)}', 'danger')
        return redirect(request.referrer)
    finally:
        cursor.close()
        conn.close()

# ==========================================
# 📺 Route: បង្ហាញអេក្រង់ទូរទស្សន៍ TV (/queue_tv)
# ==========================================
@queue_bp.route('/queue_tv')
def queue_tv():
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    
    settings = {}
    try:
        cursor.execute("SELECT setting_key, setting_value FROM queue_tv_settings")
        for row in cursor.fetchall():
            settings[row['setting_key']] = row['setting_value']
    except Exception as e:
        print(f"Error loading TV settings: {e}")
    finally:
        if cursor: cursor.close()
        if conn: conn.close()

    slide_images = []
    if os.path.exists(TV_SLIDES_FOLDER):
        valid_extensions = ('.png', '.jpg', '.jpeg', '.gif', '.webp', '.mp4', '.webm')
        slide_images = [f for f in os.listdir(TV_SLIDES_FOLDER) if f.lower().endswith(valid_extensions)]
        slide_images.sort()

    return render_template('queue_tv.html', settings=settings, slide_images=slide_images)

# ==========================================
# ⚙️ Route: ផ្ទាំងការកំណត់ទូរទស្សន៍ (TV Settings)
# ==========================================
@queue_bp.route('/queue_tv_settings', methods=['GET', 'POST'])
def manage_queue_tv_settings():
    if 'loggedin' not in session: 
        return redirect(url_for('login'))
    
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    
    if request.method == 'POST':
        action = request.form.get('action')
        
        try:
            if action in ['save_general', 'save_header', 'save_khmer', 'save_english', 'save_footer']:
                for key, val in request.form.items():
                    if key != 'action':
                        cursor.execute("INSERT INTO queue_tv_settings (setting_key, setting_value) VALUES (%s, %s) ON DUPLICATE KEY UPDATE setting_value = %s", (key, val, val))

                if action == 'save_general':
                    toggles = ['enable_chime', 'enable_voice', 'enable_video_audio']
                    for toggle in toggles:
                        val = 'true' if request.form.get(toggle) else 'false'
                        cursor.execute("INSERT INTO queue_tv_settings (setting_key, setting_value) VALUES (%s, %s) ON DUPLICATE KEY UPDATE setting_value = %s", (toggle, val, val))
                    
                    voice_speed = request.form.get('voice_speed', '1.25')
                    cursor.execute("INSERT INTO queue_tv_settings (setting_key, setting_value) VALUES ('voice_speed', %s) ON DUPLICATE KEY UPDATE setting_value = %s", (voice_speed, voice_speed))
                    
                conn.commit()
                flash('បានរក្សាទុកការកំណត់ជោគជ័យ!', 'success')

            elif action == 'upload_logo':
                file = request.files.get('tv_logo')
                if file and file.filename != '':
                    filename = secure_filename(file.filename)
                    file.save(os.path.join(current_app.root_path, 'static', 'uploads', 'profiles', filename))
                    cursor.execute("INSERT INTO queue_tv_settings (setting_key, setting_value) VALUES ('tv_logo', %s) ON DUPLICATE KEY UPDATE setting_value = %s", (filename, filename))
                    conn.commit()
                    flash('បានផ្លាស់ប្តូរឡូហ្គោថ្មីជោគជ័យ!', 'success')

            elif action == 'upload_chime':
                file = request.files.get('tv_chime')
                if file and file.filename != '':
                    filename = secure_filename(file.filename)
                    file.save(os.path.join(TV_SOUNDS_FOLDER, filename))
                    cursor.execute("INSERT INTO queue_tv_settings (setting_key, setting_value) VALUES ('chime_filename', %s) ON DUPLICATE KEY UPDATE setting_value = %s", (filename, filename))
                    conn.commit()
                    flash('បានផ្លាស់ប្តូរសំឡេងកណ្តឹងថ្មីជោគជ័យ!', 'success')

            elif action == 'upload_slide':
                file = request.files.get('slide_image')
                if file and file.filename != '':
                    filename = secure_filename(file.filename)
                    file.save(os.path.join(TV_SLIDES_FOLDER, filename))
                    flash('បានបញ្ចូលរូបភាព/វីដេអូថ្មីជោគជ័យ!', 'success')

            elif action == 'delete_slide':
                filename = request.form.get('filename')
                if filename:
                    file_path = os.path.join(TV_SLIDES_FOLDER, secure_filename(filename))
                    if os.path.exists(file_path):
                        os.remove(file_path)
                        flash('បានលុបចេញជោគជ័យ!', 'success')

        except Exception as e:
            conn.rollback()
            flash(f'មានបញ្ហា: {str(e)}', 'danger')
        
        return redirect(url_for('queue.manage_queue_tv_settings'))

    settings = {}
    try:
        cursor.execute("SELECT setting_key, setting_value FROM queue_tv_settings")
        for row in cursor.fetchall():
            settings[row['setting_key']] = row['setting_value']
    finally:
        if cursor: cursor.close()
        if conn: conn.close()

    slide_images = []
    if os.path.exists(TV_SLIDES_FOLDER):
        valid_extensions = ('.png', '.jpg', '.jpeg', '.gif', '.webp', '.mp4', '.webm')
        slide_images = [f for f in os.listdir(TV_SLIDES_FOLDER) if f.lower().endswith(valid_extensions)]
        slide_images.sort()

    return render_template('queue_tv_settings.html', settings=settings, slide_images=slide_images, full_name=session.get('full_name'))

# ==========================================
# API សម្រាប់ទូរទស្សន៍ទាញទិន្នន័យ (Auto-Refresh ស្ងាត់ៗ)
# ==========================================
@queue_bp.route('/api/queue_tv_data', methods=['GET'])
def get_queue_tv_data():
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("SELECT setting_key, setting_value FROM queue_tv_settings")
        settings = {row['setting_key']: row['setting_value'] for row in cursor.fetchall()}

        try:
            timeout_seconds = int(settings.get('auto_clear_timer', 300))
            cursor.execute("""
                UPDATE queue_tickets 
                SET status = 'Completed' 
                WHERE status = 'Processing' 
                AND TIMESTAMPDIFF(SECOND, updated_at, NOW()) > %s
            """, (timeout_seconds,))
            conn.commit()
        except:
            pass 

        cursor.execute("""
            SELECT q.queue_number, cnt.counter_name_en as counter_name, 'processing' as status, q.id 
            FROM queue_tickets q 
            LEFT JOIN counters cnt ON q.counter_no = cnt.id 
            WHERE q.status = 'Processing' 
            ORDER BY q.updated_at DESC LIMIT 5 
        """)
        processing_list = cursor.fetchall()

        limit_waiting = 5 - len(processing_list)
        waiting_list = []
        if limit_waiting > 0:
            cursor.execute("""
                SELECT queue_number, '-' as counter_name, 'waiting' as status, id 
                FROM queue_tickets 
                WHERE status = 'Waiting' 
                  AND DATE(created_at) = CURDATE() 
                ORDER BY id ASC LIMIT %s
            """, (limit_waiting,))
            waiting_list = cursor.fetchall()

        slide_images = []
        if os.path.exists(TV_SLIDES_FOLDER):
            valid_extensions = ('.png', '.jpg', '.jpeg', '.gif', '.webp', '.mp4', '.webm')
            slide_images = [f for f in os.listdir(TV_SLIDES_FOLDER) if f.lower().endswith(valid_extensions)]

        return jsonify({
            'status': 'success', 
            'tv_list': processing_list + waiting_list, 
            'settings': settings,
            'slides_count': len(slide_images)
        })
    except Exception as e:
        return jsonify({'status': 'error', 'message': str(e)})
    finally:
        if cursor: cursor.close()
        if conn: conn.close()

from datetime import datetime, timedelta
from flask import session, jsonify
# សូមប្រាកដថាលោកអ្នកបាន import datetime រួចរាល់នៅខាងលើគេនៃ File Python របស់លោកអ្នក

# ==========================================
# API សម្រាប់ទាញយកស្ថិតិជួររង់ចាំ (បង្ហាញលើផ្ទាំង Counter)
# ==========================================
@queue_bp.route('/api/queue_stats', methods=['GET'])
def get_queue_stats():
    # 🌟 ១. ការពារសុវត្ថិភាព៖ ត្រូវប្រាកដថាមន្ត្រីបាន Login រួចរាល់
    if 'loggedin' not in session:
        return jsonify({'status': 'error', 'message': 'Unauthorized'}), 401

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    
    try:
        # 🌟 ២. កំណត់ម៉ោងនៅកម្ពុជា (UTC+7) ជានិច្ច ទោះ Server នៅប្រទេសណាក៏ដោយ
        today_date_khmer = (datetime.utcnow() + timedelta(hours=7)).strftime('%Y-%m-%d')

        # ទាញយកលេខចុងក្រោយប្រចាំថ្ងៃនេះ (ប្រើប្រាស់អថេរ today_date_khmer)
        cursor.execute("""
            SELECT queue_number 
            FROM queue_tickets 
            WHERE DATE(created_at) = %s 
            ORDER BY id DESC LIMIT 1
        """, (today_date_khmer,))
        last_ticket = cursor.fetchone()
        last_number = last_ticket['queue_number'] if last_ticket else '---'

        # រាប់ចំនួនអ្នករង់ចាំប្រចាំថ្ងៃនេះ
        cursor.execute("""
            SELECT COUNT(*) AS total_waiting 
            FROM queue_tickets 
            WHERE status = 'Waiting' 
              AND DATE(created_at) = %s
        """, (today_date_khmer,))
        waiting_count = cursor.fetchone()['total_waiting']

        return jsonify({
            'status': 'success', 
            'last_number': last_number, 
            'total_waiting': waiting_count
        })
    except Exception as e:
        return jsonify({'status': 'error', 'message': str(e)})
    finally:
        if cursor: cursor.close()
        if conn: conn.close()

# ====================================================================
# ផ្ទាំងគ្រប់គ្រងការស្កេនឯកសារសម្រាប់ Admin (Admin Manage Scanners)
# ====================================================================
@queue_bp.route('/admin_manage_scanners')
def admin_manage_scanners():
    if 'loggedin' not in session:
        return redirect(url_for('login'))
        
    allowed_roles = ['Admin', 'Super Admin', 'Center President', 'Head of Reception']
    if session.get('role_name') not in allowed_roles:
        flash('សុំទោស! អ្នកមិនមានសិទ្ធិចូលទំព័រនេះទេ។', 'danger')
        return redirect(url_for('dashboard'))

    return render_template('admin_manage_scanners.html')

# ====================================================================
# API សម្រាប់ Auto-Refresh ទំព័រ Admin Manage Scanners
# ====================================================================
@queue_bp.route('/api/admin_scanners_data', methods=['GET'])
def get_admin_scanners_data():
    if 'loggedin' not in session:
        return jsonify({'status': 'error', 'message': 'Unauthorized'}), 401

    from app import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    
    try:
        # 🌟 កែតម្រូវ៖ ប្រើ DATE(q.updated_at) និងលុបការឆែក action_status ចេញដើម្បីឱ្យទិន្នន័យមកវិញ
        cursor.execute("""
            SELECT 
                COUNT(DISTINCT q.application_no) as total_approved,
                COUNT(DISTINCT CASE WHEN a.scan_status = 'Scanned' THEN q.application_no END) as total_scanned,
                COUNT(DISTINCT CASE WHEN IFNULL(a.scan_status, 'Pending') = 'Pending' THEN q.application_no END) as total_pending,
                COUNT(DISTINCT CASE WHEN a.scan_status = 'Scanned' AND c.gender IN ('F', 'FEMALE', 'ស្រី') THEN q.application_no END) as female_scanned
            FROM queue_tickets q
            JOIN candidates c ON q.application_no = c.application_no
            JOIN application_details a ON c.application_no = a.application_no
            WHERE DATE(q.updated_at) = CURDATE() AND q.status = 'Completed'
        """)
        stats = cursor.fetchone()

        cursor.execute("""
            SELECT 
                u.id, 
                u.full_name, 
                COUNT(DISTINCT q.application_no) as processed_today,
                COUNT(DISTINCT CASE WHEN c.gender IN ('F', 'FEMALE', 'ស្រី') THEN q.application_no END) as female_processed_today
            FROM users u
            JOIN roles r ON u.role_id = r.role_id
            LEFT JOIN application_details a ON u.id = a.scanned_by AND a.scan_status = 'Scanned'
            LEFT JOIN queue_tickets q ON a.application_no = q.application_no 
                AND DATE(q.updated_at) = CURDATE() 
                AND q.status = 'Completed'
            LEFT JOIN candidates c ON a.application_no = c.application_no
            WHERE r.name IN ('Scan Officer', 'Document Scanner', 'មន្ត្រីស្កេនឯកសារ') 
              AND u.is_active = 1
            GROUP BY u.id, u.full_name
            ORDER BY processed_today DESC, u.full_name ASC
        """)
        officers_list = cursor.fetchall()

        # ៣. ទាញយកបញ្ជីអ្នកដែលបាត់ខ្លួន
        cursor.execute("""
            SELECT q.queue_number, q.application_no, 
                   CONCAT(IFNULL(a.last_name_kh, ''), ' ', IFNULL(a.first_name_kh, '')) as full_name_kh, 
                   TIME_FORMAT(MIN(q.updated_at), '%H:%i') as approved_time
            FROM queue_tickets q 
            LEFT JOIN candidates c ON q.application_no = c.application_no 
            LEFT JOIN application_details a ON q.application_no = a.application_no
            WHERE DATE(q.updated_at) = CURDATE() 
              AND q.status = 'Completed' 
              AND IFNULL(a.scan_status, 'Pending') = 'Pending'
            GROUP BY q.queue_number, q.application_no, full_name_kh 
            ORDER BY approved_time ASC
        """)
        pending_list = cursor.fetchall()

        return jsonify({
            'status': 'success',
            'stats': stats,
            'officers': officers_list,
            'pending_list': pending_list
        })

    except Exception as e:
        return jsonify({'status': 'error', 'message': str(e)})
    finally:
        cursor.close()
        conn.close()

# ====================================================================
# API សម្រាប់ទាញយករបាយការណ៍ស្កេនជា Excel (តាមស្តង់ដារកូដចាស់)
# ====================================================================
from datetime import datetime
import pytz
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from flask import send_file, request, session, redirect, url_for, flash

@queue_bp.route('/export_scanners_report', methods=['GET', 'POST'])
def export_scanners_report():
    if 'loggedin' not in session:
        return redirect(url_for('login'))
        
    allowed_roles = ['Admin', 'Super Admin', 'Center President', 'Head of Reception']
    if session.get('role_name') not in allowed_roles:
        flash('សុំទោស! អ្នកមិនមានសិទ្ធិទាញយករបាយការណ៍នេះទេ។', 'danger')
        return redirect(url_for('dashboard'))

    from app import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    
    # 🌟 កំណត់ម៉ោងកម្ពុជា
    khmer_tz = pytz.timezone('Asia/Phnom_Penh')
    today_kh_date = datetime.now(khmer_tz).strftime('%Y-%m-%d')
    current_datetime = datetime.now(khmer_tz).strftime('%Y-%m-%d %H:%M:%S')
    
    # 🌟 ទទួលយកថ្ងៃខែពី Form បើគ្មាន គឺយកថ្ងៃនេះជាគោល
    start_date = request.form.get('start_date') or today_kh_date
    end_date = request.form.get('end_date') or start_date
    
    exporter_name = session.get('full_name', 'Unknown Admin')
    
    try:
        # 🌟 ប្រើ BETWEEN សម្រាប់ចន្លោះថ្ងៃ
        cursor.execute("""
            SELECT 
                u.full_name AS officer_name,
                r.name AS role_name,
                COUNT(DISTINCT q.application_no) as total_scanned,
                COUNT(DISTINCT CASE WHEN c.gender IN ('F', 'FEMALE', 'ស្រី') THEN q.application_no END) as female_scanned
            FROM users u
            JOIN roles r ON u.role_id = r.role_id
            JOIN application_details a ON u.id = a.scanned_by
            JOIN queue_tickets q ON a.application_no = q.application_no
            JOIN candidates c ON a.application_no = c.application_no
            WHERE a.scan_status = 'Scanned' 
              AND DATE(q.updated_at) BETWEEN %s AND %s 
              AND q.status = 'Completed'
            GROUP BY u.full_name, r.name
            ORDER BY total_scanned DESC
        """, (start_date, end_date))
        summary_data = cursor.fetchall()

        cursor.execute("""
            SELECT 
                u.full_name AS officer_name,
                q.queue_number, 
                c.application_no, 
                CONCAT(IFNULL(a.last_name_kh, ''), ' ', IFNULL(a.first_name_kh, '')) as candidate_name,
                c.name_en,
                c.gender,
                DATE(MIN(q.updated_at)) as date_done,
                TIME_FORMAT(MIN(q.updated_at), '%H:%i') as time_done
            FROM application_details a
            JOIN users u ON a.scanned_by = u.id
            JOIN candidates c ON a.application_no = c.application_no
            JOIN queue_tickets q ON a.application_no = q.application_no
            WHERE a.scan_status = 'Scanned' 
              AND DATE(q.updated_at) BETWEEN %s AND %s 
              AND q.status = 'Completed'
            GROUP BY u.full_name, q.queue_number, c.application_no, candidate_name, c.name_en, c.gender
            ORDER BY date_done DESC, time_done DESC
        """, (start_date, end_date))
        detail_data = cursor.fetchall()
        
        wb = openpyxl.Workbook()

        def apply_standard_styles(ws):
            font_title = Font(name='Khmer OS Siemreap', size=15, bold=True, color="0f172a")
            font_subtitle = Font(name='Khmer OS Siemreap', size=12, color="475569")
            font_header = Font(name='Khmer OS Siemreap', size=13, bold=True, color="ffffff")
            font_data = Font(name='Khmer OS Siemreap', size=11)
            
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            align_center = Alignment(horizontal='center', vertical='center', wrap_text=False)
            align_left = Alignment(horizontal='left', vertical='center', wrap_text=False)
            header_fill = PatternFill(start_color="1e293b", end_color="1e293b", fill_type="solid")

            ws.row_dimensions[1].height = 35 
            ws.row_dimensions[2].height = 25 
            ws.row_dimensions[3].height = 25
            ws.row_dimensions[4].height = 35 

            for row_idx, row in enumerate(ws.iter_rows(), start=1):
                if row_idx > 4:
                    ws.row_dimensions[row_idx].height = 28
                    
                for cell in row:
                    if row_idx == 1: 
                        cell.font = font_title
                        cell.alignment = align_center
                    elif row_idx in [2, 3]: 
                        cell.font = font_subtitle
                        cell.alignment = align_center
                    elif row_idx == 4: 
                        cell.font = font_header
                        cell.alignment = align_center
                        cell.border = thin_border
                        cell.fill = header_fill
                    else: 
                        cell.font = font_data
                        cell.border = thin_border
                        if cell.column in [1, 2, 3, 4]: 
                            cell.alignment = align_center
                        else:
                            cell.alignment = align_left

            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    if cell.row < 4: continue
                    try:
                        if cell.value:
                            cell_length = len(str(cell.value))
                            if cell_length > max_length: max_length = cell_length
                    except: pass
                
                if column == 'A':
                    ws.column_dimensions[column].width = 6 
                else:
                    adjusted_width = (max_length * 1.3) + 5
                    ws.column_dimensions[column].width = min(max(adjusted_width, 15), 60)

            max_col_letter = get_column_letter(ws.max_column)
            ws.merge_cells(f'A1:{max_col_letter}1')
            ws.merge_cells(f'A2:{max_col_letter}2')
            ws.merge_cells(f'A3:{max_col_letter}3')

        ws_summary = wb.active
        ws_summary.title = "Summary (សរុប)"
        
        # 🌟 បង្ហាញកាលបរិច្ឆេទចន្លោះថ្ងៃ
        date_display = f"{start_date} ដល់ {end_date}" if start_date != end_date else start_date
        ws_summary.append([f"របាយការណ៍សង្ខេប៖ មន្ត្រីស្កេនឯកសារ (ពីថ្ងៃទី {date_display})"])
        ws_summary.append([f"📅 ថ្ងៃទាញយក៖ {current_datetime} | 👤 ដោយ៖ {exporter_name}"])
        ws_summary.append([])
        
        headers1 = ["ល.រ", "ឈ្មោះមន្ត្រីស្កេន", "តួនាទី", "ស្កេនបានសរុប", "ក្នុងនោះមានបេក្ខជនស្រី"]
        ws_summary.append(headers1)
        
        for i, row in enumerate(summary_data, 1): 
            ws_summary.append([
                i, 
                row['officer_name'], 
                row['role_name'] or '-', 
                row['total_scanned'] or 0, 
                row['female_scanned'] or 0
            ])

        apply_standard_styles(ws_summary) 

        ws_details = wb.create_sheet(title="Detailed List (លម្អិត)")
        
        ws_details.append([f"បញ្ជីបេក្ខជនដែលបានស្កេនរួច (ពីថ្ងៃទី {date_display})"])
        ws_details.append([f"📅 ថ្ងៃទាញយក៖ {current_datetime} | 👤 ដោយ៖ {exporter_name}"])
        ws_details.append([]) 
        
        # 🌟 បន្ថែម Column ថ្ងៃខែ ព្រោះឥឡូវយើងទាញយកច្រើនថ្ងៃ
        headers2 = ["ល.រ", "កាលបរិច្ឆេទ", "ម៉ោងស្កេន", "លេខរង់ចាំ", "លេខកូដប្រឡង (EPS-ID)", "ឈ្មោះបេក្ខជន (ខ្មែរ)", "ឈ្មោះបេក្ខជន (អង់គ្លេស)", "ភេទ", "មន្ត្រីស្កេន"]
        ws_details.append(headers2)
        
        for i, row in enumerate(detail_data, 1):
            gender_kh = 'ស្រី' if row['gender'] in ('F', 'FEMALE', 'ស្រី') else 'ប្រុស'
            ws_details.append([
                i, 
                str(row['date_done']),
                row['time_done'], 
                row['queue_number'], 
                row['application_no'], 
                row['candidate_name'], 
                row['name_en'], 
                gender_kh, 
                row['officer_name']
            ])

        apply_standard_styles(ws_details) 

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"Scanner_Report_{start_date}_to_{end_date}.xlsx"
        
        return send_file(
            output, 
            as_attachment=True, 
            download_name=filename, 
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    except Exception as e:
        flash(f'មានកំហុសក្នុងការទាញយករបាយការណ៍៖ {str(e)}', 'danger')
        return redirect(url_for('queue.admin_manage_scanners'))
    finally:
        cursor.close()
        conn.close()   

# ====================================================================
# 🖥️ ផ្ទាំងប្រតិបត្តិការស្កេនឯកសារ (Scan Officer Dashboard)
# ====================================================================

@queue_bp.route('/scanner_dashboard')
def scanner_dashboard():
    if 'loggedin' not in session: return redirect(url_for('login'))
    return render_template('scanner_dashboard.html', full_name=session.get('full_name'))


from datetime import datetime, timedelta
from flask import session, request, jsonify

# ====================================================================
# ២. API ផ្ទៀងផ្ទាត់ពេលបាញ់ Barcode (Verify Candidate)
# ====================================================================
@queue_bp.route('/api/verify_candidate_scan', methods=['POST'])
def verify_candidate_scan():
    if 'loggedin' not in session: return jsonify({'status': 'error'}), 401
    app_no = request.json.get('application_no', '').strip()
    
    from app import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        # 🌟 កំណត់ម៉ោងកម្ពុជា (UTC+7)
        today_date_khmer = (datetime.utcnow() + timedelta(hours=7)).strftime('%Y-%m-%d')

        # ជំនួស CURDATE() ទៅជា %s
        cursor.execute("""
            SELECT q.id, q.status, IFNULL(a.scan_status, 'Pending') as scan_status, 
                   c.name_en, a.last_name_kh, a.first_name_kh, c.gender, q.application_no
            FROM queue_tickets q 
            LEFT JOIN candidates c ON q.application_no = c.application_no 
            LEFT JOIN application_details a ON q.application_no = a.application_no
            WHERE q.application_no = %s AND DATE(q.updated_at) = %s 
            ORDER BY q.id DESC LIMIT 1
        """, (app_no, today_date_khmer))
        candidate = cursor.fetchone()

        if not candidate: return jsonify({'status': 'error', 'message': '❌ រកមិនឃើញបេក្ខជននេះក្នុងបញ្ជីជួរថ្ងៃនេះទេ!'})
        if candidate['status'] != 'Completed': return jsonify({'status': 'error', 'message': '❌ បដិសេធ៖ មិនទាន់បញ្ចប់ការទទួលពាក្យនៅបញ្ជរឡើយ!'})
        if candidate['scan_status'] == 'Scanned': return jsonify({'status': 'error', 'message': '⚠️ បេក្ខជននេះត្រូវបានស្កេនរួចរាល់ហើយ!'})
        
        full_name_kh = f"{candidate['last_name_kh'] or ''} {candidate['first_name_kh'] or ''}".strip()
        return jsonify({
            'status': 'success', 
            'candidate': {
                'application_no': candidate['application_no'], 
                'name_en': candidate['name_en'] or 'មិនមាន', 
                'name_kh': full_name_kh or 'មិនមាន', 
                'gender': candidate['gender'] or 'U'
            }
        })
    finally:
        cursor.close()
        conn.close()


# ៣. API បញ្ជាក់ការស្កេនរួចរាល់ (Confirm Scan)
@queue_bp.route('/api/confirm_candidate_scan', methods=['POST'])
def confirm_candidate_scan():
    if 'loggedin' not in session: return jsonify({'status': 'error', 'message': 'Unauthorized'}), 401
    
    app_no = request.json.get('application_no')
    user_id = session.get('id')

    from app import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor()
    
    try:
        cursor.execute("""
            UPDATE application_details 
            SET scan_status = 'Scanned', scanned_by = %s 
            WHERE application_no = %s
        """, (user_id, app_no))
        conn.commit()
        return jsonify({'status': 'success', 'message': '✅ បានកត់ត្រាការស្កេនជោគជ័យ!'})
    except Exception as e:
        conn.rollback()
        return jsonify({'status': 'error', 'message': str(e)})
    finally:
        cursor.close(); conn.close()


from flask import jsonify

# ====================================================================
# ៤. API ទាញយកបញ្ជីរង់ចាំ (Silent Refresh សម្រាប់ផ្នែកខាងស្តាំ)
# ====================================================================
@queue_bp.route('/api/scanner_pending_list', methods=['GET'])
def get_scanner_pending_list():
    from app import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        # 🌟 កំណត់ម៉ោងកម្ពុជា (UTC+7)
        from datetime import datetime, timedelta
        today_date_khmer = (datetime.utcnow() + timedelta(hours=7)).strftime('%Y-%m-%d')

        cursor.execute("""
            SELECT 
                q.queue_number, 
                q.application_no, 
                CONCAT(IFNULL(a.last_name_kh, ''), ' ', IFNULL(a.first_name_kh, '')) as full_name_kh,
                ct.counter_name_kh AS counter_name
            FROM queue_tickets q 
            -- 🌟 ទាញយកជួរចុងក្រោយគេជានិច្ច ទោះបេក្ខជនមកចាប់លេខប៉ុន្មានដងក៏ដោយ
            JOIN (
                SELECT application_no, MAX(id) as max_id 
                FROM queue_tickets 
                WHERE DATE(IFNULL(updated_at, created_at)) = %s 
                GROUP BY application_no
            ) latest_q ON q.id = latest_q.max_id
            
            LEFT JOIN candidates c ON q.application_no = c.application_no 
            LEFT JOIN application_details a ON q.application_no = a.application_no
            LEFT JOIN counters ct ON q.counter_no = ct.id 
            
            WHERE q.status = 'Completed' 
              -- 🌟 កែតម្រូវចំនុចពិសេស៖ អនុញ្ញាតឱ្យបង្ហាញ ឱ្យតែវាមិនទាន់ 'Scanned' ឬ 'Rejected'
              AND IFNULL(a.scan_status, '') NOT IN ('Scanned', 'Rejected')
            ORDER BY q.updated_at ASC
        """, (today_date_khmer,))
        
        pending_list = cursor.fetchall()
        
        return jsonify({
            'status': 'success', 
            'pending_list': pending_list, 
            'count': len(pending_list)
        })
        
    except Exception as e:
        print(f"❌ SQL Error in Scanner Pending List: {e}")
        return jsonify({'status': 'error', 'message': str(e)})
        
    finally:
        if cursor: cursor.close()
        if conn: conn.close()

# 🌟 ៥. API សម្រាប់បោះបង់បេក្ខជនដែលខ្វះឯកសារ ឬមិនមកស្កេន (Return to Counter/Reject)
@queue_bp.route('/api/skip_candidate_scan', methods=['POST'])
def skip_candidate_scan():
    if 'loggedin' not in session: 
        return jsonify({'status': 'error', 'message': 'Unauthorized'}), 401
    
    app_no = request.json.get('application_no')
    user_id = session.get('id')

    from app import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor()
    
    try:
        conn.start_transaction()

        # ១. Update តារាង application_details 
        # 🌟 បន្ថែម status_id = 3 (Rejected) ដើម្បីឱ្យផ្ទាំង Admin មើលឃើញ
        cursor.execute("""
            UPDATE application_details 
            SET scan_status = 'Rejected', 
                status_id = 3, 
                scanned_by = %s, 
                remark = 'បដិសេធដោយមន្ត្រីស្កេន (ខ្វះឯកសារ/មិនមក)'
            WHERE application_no = %s
        """, (user_id, app_no))
        
        # ២. Update តារាង queue_tickets
        cursor.execute("""
            UPDATE queue_tickets 
            SET status = 'Rejected', 
                reject_reason = 'បដិសេធដោយមន្ត្រីស្កេន (ខ្វះឯកសារ/មិនមក)', 
                updated_at = NOW() 
            WHERE application_no = %s AND status = 'Completed'
        """, (app_no,))

        # ៣. Update តារាង candidates 
        cursor.execute("""
            UPDATE candidates 
            SET job_app_status = 'REJECTED' 
            WHERE application_no = %s
        """, (app_no,))
        
        conn.commit()
        return jsonify({'status': 'success', 'message': 'បានបដិសេធ និងទម្លាក់បេក្ខជននេះចេញពីប្រព័ន្ធដោយជោគជ័យ!'})
    except Exception as e:
        conn.rollback()
        return jsonify({'status': 'error', 'message': str(e)})
    finally:
        cursor.close()
        conn.close()

# ====================================================================
# 🌟 API សម្រាប់មុខងារ ADMIN លើផ្ទាំង Counter (Bypass ទៅ Scanner)
# ====================================================================

@queue_bp.route('/api/admin_search_rejected', methods=['POST'])
def admin_search_rejected():
    if 'loggedin' not in session or session.get('role_name') not in ['Admin', 'Super Admin']:
        return jsonify({'status': 'error', 'message': 'គ្មានសិទ្ធិ (Unauthorized)'}), 401

    search_val = request.json.get('search_val', '').strip()
    
    from datetime import datetime, timedelta
    today_date_khmer = (datetime.utcnow() + timedelta(hours=7)).strftime('%Y-%m-%d')

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        # 🌟 កែប្រែត្រង់នេះ៖ បន្ថែមលក្ខខណ្ឌឆែកមើល Status ទាំង ៣ តារាងបញ្ចូលគ្នា
        cursor.execute("""
            SELECT q.application_no, 
                   c.name_en,
                   CONCAT(IFNULL(a.last_name_kh, ''), ' ', IFNULL(a.first_name_kh, '')) AS name_kh
            FROM queue_tickets q
            LEFT JOIN candidates c ON q.application_no = c.application_no
            LEFT JOIN application_details a ON q.application_no = a.application_no
            WHERE (q.application_no = %s OR q.queue_number = %s)
              AND (
                  q.status = 'Rejected' OR 
                  c.job_app_status = 'REJECTED' OR 
                  a.status_id = 3
              )
              AND DATE(q.updated_at) = %s
            LIMIT 1
        """, (search_val, search_val, today_date_khmer))
        
        candidate = cursor.fetchone()

        if candidate:
            candidate['name_kh'] = candidate['name_kh'].strip()
            return jsonify({'status': 'success', 'candidate': candidate})
        else:
            return jsonify({'status': 'error', 'message': f'រកមិនឃើញបេក្ខជនលេខ "{search_val}" ក្នុងបញ្ជីបដិសេធថ្ងៃនេះទេ!'})
            
    except Exception as e:
        return jsonify({'status': 'error', 'message': f'Server Error: {str(e)}'})
    finally:
        cursor.close()
        conn.close()

@queue_bp.route('/api/admin_revive_to_scanner', methods=['POST'])
def admin_revive_to_scanner():
    if 'loggedin' not in session or session.get('role_name') not in ['Admin', 'Super Admin']:
        return jsonify({'status': 'error', 'message': 'គ្មានសិទ្ធិ (Unauthorized)'}), 401

    app_no = request.json.get('application_no')
    admin_name = session.get('full_name', 'Admin')

    from app import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        conn.start_transaction()

        # ១. ដូរ Status ក្នុងជួរទៅជា Completed វិញ ដើម្បីឱ្យលោតចូល Scanner
        cursor.execute("""
            UPDATE queue_tickets 
            SET status = 'Completed', reject_reason = NULL, updated_at = NOW() 
            WHERE application_no = %s AND status = 'Rejected'
        """, (app_no,))

        # ២. 🌟 កែត្រង់នេះ៖ ប្តូរ status_id = 1 (រង់ចាំ) វិញ ព្រោះគាត់ត្រូវរង់ចាំស្កេនសិន
        remark_text = f"អនុម័តឱ្យស្កេនឡើងវិញដោយ Admin: {admin_name}"
        cursor.execute("""
            UPDATE application_details 
            SET scan_status = 'Pending', status_id = 1, remark = %s 
            WHERE application_no = %s
        """, (remark_text, app_no))

        # ៣. 🌟 កែត្រង់នេះ៖ ប្តូរទៅជា PENDING វិញ
        cursor.execute("""
            UPDATE candidates 
            SET job_app_status = 'PENDING' 
            WHERE application_no = %s
        """, (app_no,))

        conn.commit()
        return jsonify({'status': 'success', 'message': f'ជោគជ័យ! លេខ {app_no} ត្រូវបានបញ្ជូនទៅបញ្ជីរង់ចាំស្កេនវិញហើយ។'})
    except Exception as e:
        conn.rollback()
        return jsonify({'status': 'error', 'message': str(e)})
    finally:
        cursor.close()
        conn.close()   
        
            